#!/usr/bin/env python3
"""
Restore the local payments table from the workbook used by import-payments.js.

This is an operational recovery helper for the local SQLite dev DB only.
It intentionally rewrites just the `payments` table and leaves the rest of the DB unchanged.
"""

import re
import sqlite3
import sys
from pathlib import Path
import os

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path(
    "/Users/helena/Library/Containers/com.kingsoft.wpsoffice.mac/Data/Library/Application Support/"
    "Kingsoft/WPS Cloud Files/userdata/qing/filecache/212320004/团队文档/捷淞/合同明细、美元交易.xlsx"
)
DB_PATH = Path("/Users/helena/Cursor/jiesong_system/backend/prisma/dev.db")


def parse_date(value):
    if value is None:
        return None
    text = str(value).strip()
    parts = text.split("/")
    if len(parts) != 3:
        return None
    month, day, year = [int(part) for part in parts]
    year = 2000 + year if year < 100 else year
    return f"{year:04d}-{month:02d}-{day:02d} 00:00:00"


def parse_amount(value):
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = (
        str(value)
        .replace(",", "")
        .replace("，", "")
        .replace(" ", "")
        .replace("$", "")
        .replace("￥", "")
    )
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def extract_exp(value):
    match = re.search(r"EXP\d{5,}", str(value or "").upper())
    return match.group(0) if match else None


def build_note(contract_ref, store, usage, year):
    normalized_ref = str(contract_ref or "").strip()
    normalized_store = str(store or "").strip()
    normalized_usage = str(usage or "").strip()
    normalized_year = str(year or "").strip()
    parts = []

    exp = extract_exp(normalized_ref)
    if exp:
        parts.append(f"合同号:{exp}")
    elif normalized_ref:
        parts.append(f"原始单号:{normalized_ref}")
    elif normalized_year:
        parts.append(f"年度:{normalized_year}")

    if normalized_store:
        parts.append(f"门店:{normalized_store}")

    if normalized_usage:
        parts.append(f"用途:{normalized_usage}")

    return " | ".join(parts)


def main():
    workbook_path = Path(
        sys.argv[1]
        if len(sys.argv) > 1
        else os.environ.get("PAYMENTS_WORKBOOK", DEFAULT_WORKBOOK)
    )
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheet = workbook["概况"]
    rows = list(sheet.iter_rows(values_only=True))

    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = OFF")
    cursor = conn.cursor()
    sales_contract_map = {
        contract_no: contract_id
        for contract_id, contract_no in cursor.execute("select id, contractNo from sales_contracts")
    }

    cursor.execute("delete from payments")

    created = 0
    income = 0.0
    expense = 0.0

    for row in rows[1:]:
        year = str(row[0] or "").strip()
        date_str = str(row[1] or "").strip()
        contract_ref = str(row[2] or "").strip()
        store = str(row[3] or "").strip()
        usage = str(row[4] or "").strip()
        amount_usd = parse_amount(row[5])

        if not date_str or amount_usd == 0:
            continue
        if not year or year == "总计":
            continue

        payment_date = parse_date(date_str)
        if not payment_date:
            continue

        is_income = usage == "收入" or amount_usd > 0
        payment_type = "INCOME" if is_income else "EXPENSE"
        abs_amount = round(abs(amount_usd), 2)
        exp = extract_exp(contract_ref)
        sales_contract_id = sales_contract_map.get(exp) if exp else None
        note = build_note(contract_ref, store, usage, year)

        cursor.execute(
            """
            insert into payments (
              id, type, purchaseContractId, salesContractId, amount, currency,
              paymentMethod, paymentDate, note, createdAt, updatedAt
            )
            values (
              lower(hex(randomblob(16))), ?, null, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now')
            )
            """,
            (
                payment_type,
                sales_contract_id,
                abs_amount,
                "USD",
                "电汇",
                payment_date,
                note,
            ),
        )

        created += 1
        if is_income:
            income += abs_amount
        else:
            expense += abs_amount

    conn.commit()
    conn.close()

    print(
        {
            "workbook": str(workbook_path),
            "created": created,
            "income": round(income, 2),
            "expense": round(expense, 2),
            "balance": round(income - expense, 2),
        }
    )


if __name__ == "__main__":
    main()
