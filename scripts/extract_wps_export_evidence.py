#!/usr/bin/env python3
"""
Input: parsed/attachment_inventory.csv + tmp/wps_11_export_list_raw/11-报关记录 + tmp/wps_12_customs_forms_raw/12-报关单
Output: parsed/evidence_extracts.csv/json + evidence_items.csv + evidence_extracts_summary.md
Pos: WPS 历史出货真实凭证正文抽取脚本；支持 11-报关记录附件、正式报关/退税 PDF 明细申报要素、出货清单/出口申报信息参考表、旧 xls 底稿与 12-报关单独立 PDF；商业发票只标记为 reference_only，不进入报关/退税写库

Note: 我被更新时，必须同步更新本头注释 + 所属目录 README/INDEX。
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from pypdf import PdfReader
except ModuleNotFoundError:  # PDF extraction remains optional for XLSX-only validation paths.
    PdfReader = None  # type: ignore[assignment]


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = REPO_ROOT / "tmp/wps_11_export_list_raw/11-报关记录"
STANDALONE_CUSTOMS_ROOT = REPO_ROOT / "tmp/wps_12_customs_forms_raw/12-报关单"
PARSED_DIR = REPO_ROOT / "tmp/wps_11_export_list_raw/parsed"
INVENTORY_PATH = PARSED_DIR / "attachment_inventory.csv"
DB_PATH = REPO_ROOT / "backend/prisma/dev.db"
PREFERRED_PACKING_PATH = PARSED_DIR / "preferred_packing_items.csv"
PREFERRED_SALES_PATH = PARSED_DIR / "preferred_sales_items.csv"
TARGET_CATEGORIES = {"customs_declaration", "export_tax_refund", "output_invoice"}
SHIPMENT_LIST_NAME_RE = re.compile(r"(出货清单|出口申报信息)")
SOURCE_NOTE_PREFIX = "[WPS_EVIDENCE]"
STANDALONE_CUSTOMS_PREFIX = "12-报关单"
REFERENCE_ONLY_EVIDENCE = {
    "2024年/1201 安纳汉姆 吴物流/一般贸易报关发票 合同 装箱单 出口报关单-1单.xls": {
        "reason": "draft_duplicate_of_formal_declaration",
        "declaration_no": "222920240004561873",
        "contract_no": "EXP2400005",
    },
    "91310000MAD74FYH58-20250604235840-出口退税用途确认发票明细..xlsx": {
        "reason": "tax_refund_invoice_list_without_declaration",
    },
}


@dataclass
class EvidenceExtract:
    relative_path: str
    category: str
    inferred_contracts: list[str]
    contract_inference: str
    suffix: str
    status: str
    extraction_method: str = ""
    text_length: int = 0
    page_count: int | None = None
    declaration_no: str | None = None
    pre_entry_no: str | None = None
    contract_no: str | None = None
    declared_at: str | None = None
    export_date: str | None = None
    bill_lading_no: str | None = None
    container_nos: list[str] = field(default_factory=list)
    package_count: float | None = None
    gross_weight: float | None = None
    net_weight: float | None = None
    invoice_nos: list[str] = field(default_factory=list)
    taxpayer_nos: list[str] = field(default_factory=list)
    item_count: int = 0
    readiness: str = "blocked"
    issues: list[str] = field(default_factory=list)
    text_preview: str = ""

    def to_row(self) -> dict[str, Any]:
        return {
            "relative_path": self.relative_path,
            "category": self.category,
            "inferred_contracts": ";".join(self.inferred_contracts),
            "contract_inference": self.contract_inference,
            "suffix": self.suffix,
            "status": self.status,
            "extraction_method": self.extraction_method,
            "text_length": self.text_length,
            "page_count": self.page_count,
            "declaration_no": self.declaration_no,
            "pre_entry_no": self.pre_entry_no,
            "contract_no": self.contract_no,
            "declared_at": self.declared_at,
            "export_date": self.export_date,
            "bill_lading_no": self.bill_lading_no,
            "container_nos": ";".join(self.container_nos),
            "package_count": self.package_count,
            "gross_weight": self.gross_weight,
            "net_weight": self.net_weight,
            "invoice_nos": ";".join(self.invoice_nos),
            "taxpayer_nos": ";".join(self.taxpayer_nos),
            "item_count": self.item_count,
            "readiness": self.readiness,
            "issues": "; ".join(self.issues),
            "text_preview": self.text_preview,
        }


def clean_text(value: str | None) -> str:
    return (value or "").replace("\u3000", " ").strip()


def normalize_extracted_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[\u0001-\u0008\u000b\u000c\u000e-\u001f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def compact_text(text: str) -> str:
    return re.sub(r"\s+", " ", text)


def parse_number(value: str | None) -> float | None:
    if not value:
        return None
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return None


def parse_date8(value: str | None) -> str | None:
    if not value or not re.fullmatch(r"\d{8}", value):
        return None
    return f"{value[0:4]}-{value[4:6]}-{value[6:8]}"


def parse_date_value(value: str | None) -> str | None:
    if not value:
        return None
    text = clean_text(value)
    if re.fullmatch(r"\d{8}", text):
        return parse_date8(text)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return text
    return None


def first_match(patterns: list[str], text: str, flags: int = 0) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text, flags)
        if match:
            return clean_text(match.group(1))
    return None


def all_unique(pattern: str, text: str, flags: int = 0) -> list[str]:
    return sorted({clean_text(match) for match in re.findall(pattern, text, flags) if clean_text(match)})


def unique_list(values: list[str]) -> list[str]:
    seen = set()
    result = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def numbers_close(left: float | None, right: float | None, tolerance: float = 0.05) -> bool:
    if left is None or right is None:
        return False
    return abs(float(left) - float(right)) <= tolerance


def compact_name(value: str | None) -> str:
    return re.sub(r"\s+", "", value or "")


def normalize_declaration_elements(value: str | None) -> str:
    text = compact_text(clean_text(value))
    if "|" not in text:
        return ""
    text = re.sub(r"\s*\|\s*", "|", text)
    text = text.replace("未 搪瓷", "未搪瓷").replace("盖 子", "盖子")
    text = re.sub(r"\s+", "", text)
    replacements = {
        "|未|搪瓷|": "|未搪瓷|",
        "|盖|子|": "|盖子|",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.strip(" |")


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def load_contract_match_sources() -> dict[str, dict[str, list[dict[str, str]]]]:
    sources: dict[str, dict[str, list[dict[str, str]]]] = {}
    for kind, path in (("packing", PREFERRED_PACKING_PATH), ("sales", PREFERRED_SALES_PATH)):
        for row in read_csv_rows(path):
            contract_no = row.get("contract_no")
            if not contract_no:
                continue
            sources.setdefault(contract_no, {"packing": [], "sales": []})[kind].append(row)
    return sources


def read_pdf_text(path: Path) -> tuple[str, int, list[str]]:
    if PdfReader is None:
        return "", 0, ["pdf_text_extractor_unavailable"]
    issues: list[str] = []
    reader = PdfReader(str(path))
    chunks = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            chunks.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - preserve per-page evidence, do not fail whole file.
            issues.append(f"page_{index}_extract_error:{exc}")
    return normalize_extracted_text("\n".join(chunks)), len(reader.pages), issues


def ocr_language_config() -> tuple[str, str]:
    result = subprocess.run(
        ["tesseract", "--list-langs"],
        check=False,
        capture_output=True,
        text=True,
    )
    languages = set(result.stdout.splitlines())
    if {"chi_sim", "eng", "snum"}.issubset(languages):
        return "chi_sim+eng+snum", "ocr_languages_chi_sim_eng_snum"
    return "eng+snum", "ocr_limited_languages_eng_snum"


def read_pdf_image_ocr_text(path: Path) -> tuple[str, int, list[str], str]:
    if not shutil.which("tesseract"):
        return "", 0, ["ocr_tesseract_not_available"], "ocr_unavailable"
    if PdfReader is None:
        return "", 0, ["pdf_image_extractor_unavailable"], "ocr_unavailable"

    ocr_languages, language_issue = ocr_language_config()
    issues = [language_issue]
    reader = PdfReader(str(path))
    chunks: list[str] = []
    image_count = 0
    with tempfile.TemporaryDirectory(prefix="wps-evidence-ocr-") as tmpdir:
        tmp_path = Path(tmpdir)
        for page_index, page in enumerate(reader.pages, start=1):
            for image_index, image in enumerate(page.images, start=1):
                image_count += 1
                suffix = Path(image.name).suffix.lower() or ".png"
                image_path = tmp_path / f"page-{page_index}-image-{image_index}{suffix}"
                try:
                    image_path.write_bytes(image.data)
                    result = subprocess.run(
                        ["tesseract", str(image_path), "stdout", "-l", ocr_languages, "--psm", "6"],
                        check=False,
                        capture_output=True,
                        text=True,
                    )
                except Exception as exc:  # noqa: BLE001 - OCR is best-effort evidence extraction.
                    issues.append(f"ocr_page_{page_index}_image_{image_index}_error:{exc}")
                    continue
                if result.returncode != 0:
                    stderr = compact_text(result.stderr)[:160]
                    issues.append(f"ocr_page_{page_index}_image_{image_index}_error:{stderr}")
                    continue
                text = clean_text(result.stdout)
                if text:
                    chunks.append(text)

    return normalize_extracted_text("\n".join(chunks)), image_count, unique_list(issues), f"ocr_image_{ocr_languages.replace('+', '_')}"


def read_excel_text(path: Path) -> tuple[str, list[str]]:
    if path.suffix.lower() == ".xls":
        return read_legacy_xls_text(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        chunks.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [clean_text(str(value)) for value in row if value is not None and clean_text(str(value))]
            if values:
                chunks.append(" | ".join(values))
    text = normalize_extracted_text("\n".join(chunks))
    if path.suffix.lower() == ".xlsx" and len(text) < 200:
        fallback_text, fallback_issues = read_xlsx_text_via_frontend(path)
        if len(fallback_text) > len(text):
            return fallback_text, ["xlsx_text_via_frontend_xlsx_fallback", *fallback_issues]
    return text, []


def read_legacy_xls_text(path: Path) -> tuple[str, list[str]]:
    text, issues = read_xlsx_text_via_frontend(path)
    return text, ["legacy_xls_read_via_frontend_xlsx", *issues]


def read_xlsx_text_via_frontend(path: Path) -> tuple[str, list[str]]:
    xlsx_module = REPO_ROOT / "frontend/node_modules/xlsx"
    if not xlsx_module.exists():
        return "", ["frontend_xlsx_dependency_missing"]
    node_bin = shutil.which("node")
    if not node_bin:
        return "", ["frontend_xlsx_requires_node"]

    script = """
const XLSX = require(process.argv[1]);
const workbook = XLSX.readFile(process.argv[2], { cellDates: false, raw: false });
const chunks = [];
for (const name of workbook.SheetNames) {
  chunks.push(`## ${name}`);
  const rows = XLSX.utils.sheet_to_json(workbook.Sheets[name], {
    header: 1,
    blankrows: false,
    defval: '',
  });
  for (const row of rows) {
    const values = row.map((value) => String(value || '').trim()).filter(Boolean);
    if (values.length > 0) chunks.push(values.join(' | '));
  }
}
process.stdout.write(JSON.stringify({ text: chunks.join('\\n') }));
"""
    result = subprocess.run(
        [node_bin, "-e", script, str(xlsx_module), str(path)],
        check=False,
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        return "", [f"frontend_xlsx_read_error:{compact_text(result.stderr)[:160]}"]
    try:
        payload = json.loads(result.stdout or "{}")
    except json.JSONDecodeError as exc:
        return "", [f"frontend_xlsx_json_error:{exc}"]
    return normalize_extracted_text(str(payload.get("text") or "")), []


def extract_invoice_nos(text: str) -> list[str]:
    candidates = all_unique(r"\b\d{20}\b", text)
    return [item for item in candidates if not item.startswith(("5304", "5316", "2229"))]


def is_commercial_invoice_text(text: str) -> bool:
    one_line = compact_text(text).upper()
    return "COMMERCIAL INVOICE" in one_line


def extract_taxpayer_nos(text: str) -> list[str]:
    return all_unique(r"\b[0-9A-Z]{15,20}\b", text)


def extract_declaration_elements_by_item(text: str) -> dict[tuple[int, str, str], str]:
    result: dict[tuple[int, str, str], str] = {}
    item_with_elements_pattern = re.compile(
        r"(?:^|\s)(\d{1,3})\s+(\d{8,13})\s+(.{1,80}?)\s+法定数量/法定单位\s+"
        r"[0-9.]+\s*[\u4e00-\u9fffA-Za-z]+\s+美国\s+[0-9.]+\s+[0-9.]+\s+USD\s+"
        r"(.{1,200}?)\s+第二数量/第二单位",
    )
    for match in item_with_elements_pattern.finditer(text):
        item_no, hs_code, name, declaration_elements = match.groups()
        normalized_elements = normalize_declaration_elements(declaration_elements)
        if not normalized_elements:
            continue
        clean_name = re.sub(r"\s+", "", name)[:80]
        result[(int(item_no), hs_code, clean_name)] = normalized_elements
    return result


def extract_items(text: str, evidence: EvidenceExtract) -> list[dict[str, Any]]:
    one_line = compact_text(text)
    declaration_elements_by_item = extract_declaration_elements_by_item(one_line)
    items: list[dict[str, Any]] = []
    item_pattern = re.compile(
        r"(?:^|\s)(\d{1,3})\s+(\d{8,13})\s+(.{1,80}?)\s+法定数量/法定单位\s+"
        r"([0-9.]+)\s*([\u4e00-\u9fffA-Za-z]+)\s+美国\s+([0-9.]+)\s+([0-9.]+)\s+USD",
    )
    for match in item_pattern.finditer(one_line):
        item_no, hs_code, name, quantity, unit, unit_price, total_price = match.groups()
        clean_name = re.sub(r"\s+", "", name)
        declaration_elements = declaration_elements_by_item.get((int(item_no), hs_code, clean_name[:80]), "")
        items.append({
            "relative_path": evidence.relative_path,
            "category": evidence.category,
            "inferred_contracts": ";".join(evidence.inferred_contracts),
            "declaration_no": evidence.declaration_no,
            "contract_no": evidence.contract_no,
            "item_no": int(item_no),
            "hs_code": hs_code,
            "customs_name": clean_name[:80],
            "quantity": parse_number(quantity),
            "unit": unit,
            "unit_price": parse_number(unit_price),
            "total_price": parse_number(total_price),
            "currency": "USD",
            "declaration_elements": declaration_elements,
            "source_note": f"{SOURCE_NOTE_PREFIX} {evidence.relative_path}",
        })
    customs_pdf_pattern = re.compile(
        r"(?:^|\s)(\d{1,3})\s+(\d{8,13})\s+(.{1,80}?)\s+"
        r"([0-9.]+)\s*([\u4e00-\u9fffA-Za-z]+)\s+([0-9.]+)\s+中国\s+美国.*?"
        r"\s([0-9.]+)\s+\(CHN\)",
    )
    existing_keys = {(item["item_no"], item["hs_code"], item["customs_name"]) for item in items}
    for match in customs_pdf_pattern.finditer(one_line):
        item_no, hs_code, name, quantity, unit, unit_price, total_price = match.groups()
        clean_name = re.sub(r"\s+", "", name)
        key = (int(item_no), hs_code, clean_name[:80])
        if key in existing_keys:
            continue
        declaration_elements = declaration_elements_by_item.get(key, "")
        items.append({
            "relative_path": evidence.relative_path,
            "category": evidence.category,
            "inferred_contracts": ";".join(evidence.inferred_contracts),
            "declaration_no": evidence.declaration_no,
            "contract_no": evidence.contract_no,
            "item_no": int(item_no),
            "hs_code": hs_code,
            "customs_name": clean_name[:80],
            "quantity": parse_number(quantity),
            "unit": unit,
            "unit_price": parse_number(unit_price),
            "total_price": parse_number(total_price),
            "currency": "USD",
            "declaration_elements": declaration_elements,
            "source_note": f"{SOURCE_NOTE_PREFIX} {evidence.relative_path}",
        })
    return items


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", clean_text(str(value) if value is not None else ""))


def column_value(row: tuple[Any, ...], columns: dict[str, int], key: str) -> str:
    index = columns.get(key)
    if index is None or index >= len(row):
        return ""
    value = row[index]
    return clean_text(str(value) if value is not None else "")


def parse_item_no(value: str) -> int | None:
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        return None


def extract_shipment_list_items(path: Path, evidence: EvidenceExtract) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    items: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        header_columns: dict[str, int] | None = None
        for row in worksheet.iter_rows(values_only=True):
            normalized = [normalize_header(value) for value in row]
            if header_columns is None:
                if "HSCODE" not in normalized or "申报信息" not in normalized:
                    continue
                header_columns = {}
                for index, header in enumerate(normalized):
                    if header in {"Item", "序号"}:
                        header_columns["item_no"] = index
                    elif header in {"货物名称及规格NameofCommodityandSpecifications", "货物名称及规格", "NameofCommodityandSpecifications"}:
                        header_columns["customs_name"] = index
                    elif header == "HSCODE":
                        header_columns["hs_code"] = index
                    elif header == "申报信息":
                        header_columns["declaration_elements"] = index
                    elif header == "境内货源地":
                        header_columns["source_area"] = index
                    elif header == "数量":
                        header_columns["quantity"] = index
                    elif header == "单位":
                        header_columns["unit"] = index
                continue
            item_no = parse_item_no(column_value(row, header_columns, "item_no"))
            hs_code = column_value(row, header_columns, "hs_code")
            customs_name = column_value(row, header_columns, "customs_name")
            declaration_elements = column_value(row, header_columns, "declaration_elements")
            if item_no is None or not hs_code or not customs_name or not declaration_elements:
                continue
            items.append({
                "relative_path": evidence.relative_path,
                "category": evidence.category,
                "inferred_contracts": ";".join(evidence.inferred_contracts),
                "declaration_no": evidence.declaration_no,
                "contract_no": evidence.contract_no,
                "item_no": item_no,
                "hs_code": hs_code,
                "customs_name": customs_name[:80],
                "quantity": parse_number(column_value(row, header_columns, "quantity")),
                "unit": column_value(row, header_columns, "unit"),
                "unit_price": None,
                "total_price": None,
                "currency": "",
                "declaration_elements": declaration_elements,
                "source_area": column_value(row, header_columns, "source_area"),
                "source_note": f"{SOURCE_NOTE_PREFIX} {evidence.relative_path}",
            })
    return items


def score_contract_match(
    evidence: EvidenceExtract,
    items: list[dict[str, Any]],
    rows: dict[str, list[dict[str, str]]],
) -> int:
    score = 0
    for item in items:
        item_name = compact_name(item.get("customs_name"))
        for row in rows.get("sales", []):
            row_name = compact_name(row.get("product_name"))
            if item_name and row_name and (item_name in row_name or row_name in item_name):
                score += 1
            if numbers_close(item.get("quantity"), parse_number(row.get("quantity"))):
                score += 3
            if numbers_close(item.get("unit_price"), parse_number(row.get("unit_price")), tolerance=0.02):
                score += 2
            if numbers_close(item.get("total_price"), parse_number(row.get("total_price")), tolerance=0.1):
                score += 4
        for row in rows.get("packing", []):
            row_name = compact_name(row.get("product_name"))
            if item_name and row_name and (item_name in row_name or row_name in item_name):
                score += 1
            if numbers_close(item.get("quantity"), parse_number(row.get("quantity"))):
                score += 3
            if numbers_close(evidence.gross_weight, parse_number(row.get("gross_weight"))):
                score += 2
            if numbers_close(evidence.net_weight, parse_number(row.get("net_weight"))):
                score += 2
            if evidence.package_count is not None and numbers_close(evidence.package_count, parse_number(row.get("boxes"))):
                score += 1
    return score


def infer_contract_by_item_match(evidence: EvidenceExtract, items: list[dict[str, Any]]) -> None:
    if evidence.contract_no or len(evidence.inferred_contracts) <= 1 or not items:
        return
    sources = load_contract_match_sources()
    scores = {
        contract_no: score_contract_match(evidence, items, sources.get(contract_no, {}))
        for contract_no in evidence.inferred_contracts
    }
    best = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    if len(best) < 2 or best[0][1] < 5 or best[0][1] == best[1][1]:
        return
    evidence.contract_no = best[0][0]
    evidence.contract_inference = f"item_match:{evidence.contract_inference}"


def fill_fields(evidence: EvidenceExtract, text: str) -> None:
    one_line = compact_text(text)
    evidence.declaration_no = first_match([
        r"海关编号[：:\s|]*([0-9]{18})",
        r"\*([0-9]{18})\*",
        r"\b(5[0-9]{17})\b",
    ], one_line)
    evidence.pre_entry_no = first_match([
        r"预录入编号[：:\s|]*([A-Z][0-9]{8,})",
    ], one_line)
    evidence.contract_no = first_match([
        r"合同协议号[：:\s|]*([A-Z]{3}[0-9]{6,8})",
        r"\b(EXP[0-9]{6,8})\b",
    ], one_line, flags=re.IGNORECASE)
    if evidence.contract_no:
        evidence.contract_no = evidence.contract_no.upper()
    evidence.declared_at = parse_date_value(first_match([r"申报日期[：:\s|]*([0-9]{8})"], one_line))
    evidence.export_date = parse_date_value(first_match([r"出口日期[：:\s|]*([0-9]{8})"], one_line))
    if not evidence.export_date or not evidence.declared_at:
        date_match = re.search(r"出口日期.*?申报日期.*?([0-9]{8})\s+([0-9]{4}-[0-9]{2}-[0-9]{2})", one_line)
        if date_match:
            evidence.export_date = evidence.export_date or parse_date_value(date_match.group(1))
            evidence.declared_at = evidence.declared_at or parse_date_value(date_match.group(2))
    evidence.bill_lading_no = first_match([
        r"提运单号[：:\s|]*([A-Z0-9/\-]{6,40})",
    ], one_line)
    evidence.container_nos = all_unique(r"\b[A-Z]{4}[0-9]{7}\b", one_line)
    evidence.package_count = parse_number(first_match([r"件数[：:\s|]*([0-9.]+)"], one_line))
    evidence.gross_weight = parse_number(first_match([r"毛重（千克）[：:\s|]*([0-9.]+)", r"毛重\(千克\)[：:\s|]*([0-9.]+)"], one_line))
    evidence.net_weight = parse_number(first_match([r"净重（千克）[：:\s|]*([0-9.]+)", r"净重\(千克\)[：:\s|]*([0-9.]+)"], one_line))
    if evidence.package_count is None or evidence.gross_weight is None or evidence.net_weight is None:
        weight_match = re.search(r"包装种类.*?件数.*?毛重.*?净重.*?([0-9]+)\s+([0-9.]+)\s+([0-9.]+)\s+FOB", one_line)
        if weight_match:
            evidence.package_count = evidence.package_count or parse_number(weight_match.group(1))
            evidence.gross_weight = evidence.gross_weight or parse_number(weight_match.group(2))
            evidence.net_weight = evidence.net_weight or parse_number(weight_match.group(3))
    evidence.invoice_nos = extract_invoice_nos(one_line)
    evidence.taxpayer_nos = extract_taxpayer_nos(one_line)


def unique_db_declaration_for_contract(contract_no: str | None) -> str | None:
    if not contract_no:
        return None
    conn = sqlite3.connect(DB_PATH)
    rows = conn.execute(
        """
        select cd.declarationNo
        from customs_declarations cd
        join sales_contracts sc on sc.id = cd.salesContractId
        where sc.contractNo = ?
        order by cd.declarationNo
        """,
        (contract_no,),
    ).fetchall()
    conn.close()
    if len(rows) != 1:
        return None
    return rows[0][0]


def infer_declaration_from_unique_contract(evidence: EvidenceExtract) -> None:
    if evidence.category not in {"export_tax_refund", "shipment_list"} or evidence.declaration_no:
        return
    contract_no = evidence.contract_no
    if not contract_no and len(evidence.inferred_contracts) == 1:
        contract_no = evidence.inferred_contracts[0]
        if evidence.category == "shipment_list":
            evidence.contract_no = contract_no
    declaration_no = unique_db_declaration_for_contract(contract_no)
    if not declaration_no:
        return
    evidence.declaration_no = declaration_no
    evidence.contract_inference = f"unique_contract_customs:{evidence.contract_inference}"


def assess_readiness(evidence: EvidenceExtract) -> None:
    inferred = set(evidence.inferred_contracts)
    if evidence.status != "ok":
        evidence.readiness = "blocked"
        return
    if any(
        issue.startswith(("draft_duplicate_of_formal_declaration", "tax_refund_invoice_list_without_declaration"))
        for issue in evidence.issues
    ):
        evidence.readiness = "reference_only"
        return
    if evidence.category == "shipment_list":
        if not evidence.declaration_no:
            evidence.issues.append("missing_unique_declaration_for_shipment_list")
        if evidence.item_count == 0:
            evidence.issues.append("missing_structured_shipment_items")
        evidence.readiness = "reference_only" if not evidence.issues else "needs_review"
        return
    if not evidence.declaration_no and evidence.category in {"customs_declaration", "export_tax_refund"}:
        evidence.issues.append("missing_declaration_no")
    resolved_by_contract_no = bool(evidence.contract_no and (not inferred or evidence.contract_no in inferred))
    if len(inferred) != 1 and not resolved_by_contract_no:
        evidence.issues.append("ambiguous_or_missing_contract_inference")
    if evidence.contract_no and inferred and evidence.contract_no not in inferred:
        evidence.issues.append("contract_no_mismatch")
    if evidence.category == "output_invoice" and "commercial_invoice_only" in evidence.issues:
        evidence.readiness = "reference_only"
        return
    if evidence.category == "output_invoice" and not evidence.invoice_nos:
        evidence.issues.append("missing_invoice_no")

    if evidence.issues:
        evidence.readiness = "needs_review"
    elif (
        evidence.category in {"customs_declaration", "export_tax_refund"}
        and evidence.declaration_no
        and (len(inferred) == 1 or resolved_by_contract_no)
    ):
        evidence.readiness = "ready_for_mapping"
    elif evidence.category == "output_invoice" and evidence.invoice_nos and len(inferred) == 1:
        evidence.readiness = "ready_for_mapping"
    else:
        evidence.readiness = "needs_review"


def is_draft_duplicate_reference(evidence: EvidenceExtract) -> bool:
    return any(issue.startswith("draft_duplicate_of_formal_declaration") for issue in evidence.issues)


def is_tax_refund_invoice_list_reference(evidence: EvidenceExtract) -> bool:
    return any(issue.startswith("tax_refund_invoice_list_without_declaration") for issue in evidence.issues)


def load_inventory() -> list[dict[str, str]]:
    with INVENTORY_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = []
        for row in csv.DictReader(handle):
            if row["category"] in TARGET_CATEGORIES:
                rows.append(row)
            elif row["category"] == "spreadsheet_other" and SHIPMENT_LIST_NAME_RE.search(row["relative_path"]):
                shipment_row = dict(row)
                shipment_row["category"] = "shipment_list"
                rows.append(shipment_row)
    if STANDALONE_CUSTOMS_ROOT.exists():
        for path in sorted(STANDALONE_CUSTOMS_ROOT.glob("*.pdf")):
            contract_match = re.search(r"\b(EXP[0-9]{6,8})\b", path.name, flags=re.IGNORECASE)
            rows.append({
                "relative_path": f"{STANDALONE_CUSTOMS_PREFIX}/{path.name}",
                "category": "customs_declaration",
                "inferred_contracts": contract_match.group(1).upper() if contract_match else "",
                "contract_inference": "standalone_customs_filename",
                "suffix": path.suffix.lower(),
            })
    return rows


def resolve_evidence_path(relative_path: str) -> Path:
    if relative_path.startswith(f"{STANDALONE_CUSTOMS_PREFIX}/"):
        return STANDALONE_CUSTOMS_ROOT / relative_path[len(STANDALONE_CUSTOMS_PREFIX) + 1:]
    return SOURCE_ROOT / relative_path


def extract_one(row: dict[str, str]) -> tuple[EvidenceExtract, list[dict[str, Any]]]:
    relative_path = row["relative_path"]
    path = resolve_evidence_path(relative_path)
    evidence = EvidenceExtract(
        relative_path=relative_path,
        category=row["category"],
        inferred_contracts=[item for item in row["inferred_contracts"].split(";") if item],
        contract_inference=row["contract_inference"],
        suffix=row["suffix"],
        status="ok",
    )

    try:
        if path.suffix.lower() == ".pdf":
            text, pages, issues = read_pdf_text(path)
            evidence.page_count = pages
            evidence.extraction_method = "pdf_text"
            evidence.issues.extend(issues)
            if not text:
                ocr_text, ocr_image_count, ocr_issues, ocr_method = read_pdf_image_ocr_text(path)
                evidence.issues.extend(ocr_issues)
                if ocr_text:
                    text = ocr_text
                    evidence.extraction_method = ocr_method
                elif ocr_image_count > 0:
                    evidence.extraction_method = f"{ocr_method}_empty"
        elif path.suffix.lower() in {".xlsx", ".xls"}:
            text, issues = read_excel_text(path)
            evidence.extraction_method = "xlsx_text" if path.suffix.lower() == ".xlsx" else "legacy_xls_text"
            evidence.issues.extend(issues)
        else:
            text = ""
            evidence.extraction_method = "unsupported_suffix"
            evidence.issues.append("unsupported_suffix")
    except Exception as exc:  # noqa: BLE001 - extraction report must preserve failed files.
        evidence.status = "parse_error"
        evidence.issues.append(str(exc))
        text = ""

    evidence.text_length = len(text)
    evidence.text_preview = compact_text(text)[:500]
    if evidence.status == "ok" and evidence.text_length == 0:
        evidence.status = "empty_text"
        evidence.issues.append("no_extractable_text")

    if evidence.status == "ok":
        fill_fields(evidence, text)
        if evidence.category == "output_invoice" and is_commercial_invoice_text(text):
            evidence.issues.append("commercial_invoice_only")
        if relative_path in REFERENCE_ONLY_EVIDENCE:
            reference = REFERENCE_ONLY_EVIDENCE[relative_path]
            evidence.declaration_no = reference.get("declaration_no")
            evidence.contract_no = reference.get("contract_no")
            reference_detail = reference.get("declaration_no") or reference.get("contract_no") or "source_only"
            evidence.issues.append(f"{reference['reason']}:{reference_detail}")
        infer_declaration_from_unique_contract(evidence)
        if evidence.category == "shipment_list":
            items = extract_shipment_list_items(path, evidence)
        else:
            items = extract_items(text, evidence)
        infer_contract_by_item_match(evidence, items)
        if evidence.contract_no:
            for item in items:
                item["contract_no"] = evidence.contract_no
        if evidence.declaration_no:
            for item in items:
                item["declaration_no"] = evidence.declaration_no
        evidence.item_count = len(items)
    else:
        items = []

    assess_readiness(evidence)
    return evidence, items


def write_outputs(extracts: list[EvidenceExtract], items: list[dict[str, Any]]) -> None:
    PARSED_DIR.mkdir(parents=True, exist_ok=True)
    rows = [extract.to_row() for extract in extracts]
    csv_path = PARSED_DIR / "evidence_extracts.csv"
    json_path = PARSED_DIR / "evidence_extracts.json"
    items_path = PARSED_DIR / "evidence_items.csv"
    mapping_path = PARSED_DIR / "evidence_db_mapping.csv"
    summary_path = PARSED_DIR / "evidence_extracts_summary.md"
    mapping_rows = build_db_mapping(extracts, items)

    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()) if rows else [])
        if rows:
            writer.writeheader()
            writer.writerows(rows)

    item_fields = [
        "relative_path",
        "category",
        "inferred_contracts",
        "declaration_no",
        "contract_no",
        "item_no",
        "hs_code",
        "customs_name",
        "quantity",
        "unit",
        "unit_price",
        "total_price",
        "currency",
        "declaration_elements",
        "source_area",
        "source_note",
    ]
    with items_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=item_fields)
        writer.writeheader()
        writer.writerows(items)

    mapping_fields = [
        "relative_path",
        "category",
        "readiness",
        "inferred_contracts",
        "contract_no",
        "matched_contract_no",
        "sales_contract_id",
        "declaration_no",
        "existing_customs_declaration_id",
        "existing_tax_refund_id",
        "item_count",
        "item_total_amount",
        "mapping_action",
        "mapping_issues",
    ]
    with mapping_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=mapping_fields)
        writer.writeheader()
        writer.writerows(mapping_rows)

    summary = {
        "extract_count": len(extracts),
        "item_count": len(items),
        "status_counts": dict(sorted(Counter(item.status for item in extracts).items())),
        "category_counts": dict(sorted(Counter(item.category for item in extracts).items())),
        "readiness_counts": dict(sorted(Counter(item.readiness for item in extracts).items())),
        "mapping_action_counts": dict(sorted(Counter(row["mapping_action"] for row in mapping_rows).items())),
        "ready_for_mapping": [
            extract.to_row()
            for extract in extracts
            if extract.readiness == "ready_for_mapping"
        ],
    }
    json_path.write_text(
        json.dumps({"summary": summary, "extracts": rows, "items": items}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    lines = [
        "# WPS 真实凭证正文抽取",
        "",
        f"- 抽取文件数：`{summary['extract_count']}`",
        f"- 抽取明细行数：`{summary['item_count']}`",
        "",
        "## 状态计数",
        "",
    ]
    for key, count in summary["status_counts"].items():
        lines.append(f"- `{key}`: `{count}`")
    lines.extend(["", "## Readiness", ""])
    for key, count in summary["readiness_counts"].items():
        lines.append(f"- `{key}`: `{count}`")
    lines.extend(["", "## 数据库映射动作", ""])
    for key, count in summary["mapping_action_counts"].items():
        lines.append(f"- `{key}`: `{count}`")
    lines.extend(["", "## 可进入映射复核的文件", ""])
    for row in summary["ready_for_mapping"]:
        lines.append(
            f"- `{row['category']}` `{row['inferred_contracts']}` "
            f"`{row['declaration_no'] or row['invoice_nos']}` {row['relative_path']}"
        )
    summary_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def load_db_maps() -> dict[str, dict[str, dict[str, Any]]]:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    contracts = {
        row["contractNo"]: dict(row)
        for row in conn.execute('select id, contractNo from sales_contracts')
    }
    declarations = {
        row["declarationNo"]: dict(row)
        for row in conn.execute('select id, declarationNo, salesContractId from customs_declarations')
    }
    refunds = {
        row["customsDeclarationId"]: dict(row)
        for row in conn.execute('select id, customsDeclarationId, refundNo from tax_refunds')
    }
    return {
        "contracts": contracts,
        "declarations": declarations,
        "refunds": refunds,
    }


def build_db_mapping(extracts: list[EvidenceExtract], items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    db = load_db_maps()
    items_by_source = Counter()
    totals_by_source: dict[str, float] = {}
    for item in items:
        key = item["relative_path"]
        items_by_source[key] += 1
        totals_by_source[key] = totals_by_source.get(key, 0.0) + float(item.get("total_price") or 0)

    rows = []
    for extract in extracts:
        issues = []
        matched_contract_no = extract.contract_no
        if not matched_contract_no and len(extract.inferred_contracts) == 1:
            matched_contract_no = extract.inferred_contracts[0]
        contract = db["contracts"].get(matched_contract_no or "")
        declaration = db["declarations"].get(extract.declaration_no or "")
        refund = db["refunds"].get(declaration["id"]) if declaration else None

        if not contract:
            issues.append("missing_sales_contract")
        if extract.readiness not in {"ready_for_mapping", "reference_only"}:
            issues.append(f"extract_{extract.readiness}")
        if extract.category in {"customs_declaration", "export_tax_refund"} and not extract.declaration_no:
            issues.append("missing_declaration_no")
        if extract.readiness == "reference_only" and is_draft_duplicate_reference(extract):
            action = "evidence_reference_only"
        elif extract.readiness == "reference_only" and is_tax_refund_invoice_list_reference(extract):
            action = "tax_refund_invoice_list_reference_only"
        elif extract.category == "customs_declaration":
            if declaration:
                action = "update_existing_customs_declaration"
            elif contract and extract.readiness == "ready_for_mapping":
                action = "create_customs_declaration_draft"
            else:
                action = "blocked"
        elif extract.category == "export_tax_refund":
            if declaration and refund:
                action = "update_existing_tax_refund"
            elif declaration and contract and extract.readiness == "ready_for_mapping":
                action = "create_tax_refund_draft"
            elif contract and extract.declaration_no and extract.readiness == "ready_for_mapping":
                action = "needs_customs_declaration_first"
            else:
                action = "blocked"
        elif extract.category == "output_invoice":
            if extract.readiness == "reference_only":
                action = "commercial_invoice_reference_only"
            elif contract and extract.invoice_nos and extract.readiness == "ready_for_mapping":
                action = "map_output_invoice_to_contract"
            else:
                action = "blocked"
        elif extract.category == "shipment_list":
            if declaration and contract and extract.item_count > 0:
                action = "map_shipment_list_declaration_elements"
            else:
                action = "blocked"
        else:
            action = "blocked"

        rows.append({
            "relative_path": extract.relative_path,
            "category": extract.category,
            "readiness": extract.readiness,
            "inferred_contracts": ";".join(extract.inferred_contracts),
            "contract_no": extract.contract_no,
            "matched_contract_no": matched_contract_no,
            "sales_contract_id": contract["id"] if contract else None,
            "declaration_no": extract.declaration_no,
            "existing_customs_declaration_id": declaration["id"] if declaration else None,
            "existing_tax_refund_id": refund["id"] if refund else None,
            "item_count": items_by_source[extract.relative_path],
            "item_total_amount": round(totals_by_source.get(extract.relative_path, 0.0), 4),
            "mapping_action": action,
            "mapping_issues": "; ".join(unique_list(issues + extract.issues)),
        })
    return rows


def main() -> int:
    if not INVENTORY_PATH.exists():
        print(f"missing inventory: {INVENTORY_PATH}", file=sys.stderr)
        return 1
    extracts: list[EvidenceExtract] = []
    items: list[dict[str, Any]] = []
    for row in load_inventory():
        evidence, row_items = extract_one(row)
        extracts.append(evidence)
        items.extend(row_items)
    write_outputs(extracts, items)
    summary = {
        "extract_count": len(extracts),
        "item_count": len(items),
        "status_counts": dict(sorted(Counter(item.status for item in extracts).items())),
        "readiness_counts": dict(sorted(Counter(item.readiness for item in extracts).items())),
        "mapping_action_counts": dict(sorted(Counter(row["mapping_action"] for row in build_db_mapping(extracts, items)).items())),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
