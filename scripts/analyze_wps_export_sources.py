#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Input: tmp/wps_11_export_list_raw/11-报关记录 下的 WPS 出货/出口合同 Excel 文件
Output: tmp/wps_11_export_list_raw/parsed/ 下的合同、装箱、发票汇总与现库差异报告
Pos: WPS 出货源文件解析与导入前核对脚本；只读源文件和数据库，不直接落库；混合门店汇总行会被可加总的拆分行替代

Note: 我被更新时，必须同步更新本头注释 + 所属目录 README/INDEX。
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sqlite3
from collections import defaultdict
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


EXP_RE = re.compile(r"EXP\s*[-_]?\s*(\d{6,7})(?!\d)", re.IGNORECASE)
CG_RE = re.compile(r"CG\s*[-_]?\s*(\d{7,8})(?!\d)", re.IGNORECASE)
INVALID_FILENAME_MARKERS = ("替换模板", "template", "~$")


PACKING_HEADER_ALIASES = {
    "item_no": ("序号", "item"),
    "product_name": ("报关名", "货物名称", "商品名称", "品名", "product", "description"),
    "supplement": ("商品补充信息", "申报明细", "申报要素"),
    "store": ("门店", "发货店铺", "店铺"),
    "port": ("港口", "目的港"),
    "quantity": ("报关数量", "数量", "qty", "quantity"),
    "unit": ("单位", "unit"),
    "manufacturer": ("厂家", "供应商", "生产厂家"),
    "specification": ("规格", "尺寸", "包装/规格", "型号"),
    "boxes": ("箱数", "数量（箱）", "数量(箱)", "ctns", "carton"),
    "gross_weight": ("毛重", "gross"),
    "net_weight": ("净重", "net"),
    "volume": ("体积", "cbm", "volume"),
    "container_label": ("柜子编号", "货柜编号", "container"),
    "shipped_at": ("出货日期", "装柜日期", "发货日期"),
    "contract_no": ("合同号", "出口合同号"),
    "customs_broker": ("报关公司",),
    "is_fumigated": ("是否熏蒸", "熏蒸"),
    "purchase_contract_no": ("购销合同号", "采购合同号"),
    "purchase_cost": ("采购金额", "采购价"),
    "invoice_no": ("发票号码", "发票号"),
    "note": ("备注", "note"),
    "hs_code": ("hscode", "hs code", "hs编码", "hs"),
    "unit_price": ("单价", "unit price"),
    "total_price": ("金额", "total amount", "amount"),
}

SALES_HEADER_ALIASES = {
    "item_no": ("序号", "item"),
    "product_name": ("商品名称", "货物名称", "品名", "description"),
    "specification": ("包装/规格", "包装", "package", "规格", "尺寸"),
    "quantity": ("数量", "quantity", "qty"),
    "unit": ("单位", "unit"),
    "unit_price": ("单价", "unit price"),
    "total_price": ("金额", "total amount", "amount"),
}

INVOICE_HEADER_ALIASES = {
    "item_no": ("序号",),
    "product_name": ("产品", "商品名称", "货物名称"),
    "unit": ("单位",),
    "quantity": ("数量",),
    "unit_price_without_tax": ("单价（不含税）", "单价(不含税)", "不含税单价"),
    "amount_without_tax": ("金额", "不含税金额"),
    "tax_rate": ("税率",),
    "tax_amount": ("税额",),
    "amount_with_tax": ("总金额", "价税合计"),
    "declaration_no": ("报关单号",),
    "invoice_no": ("发票号码", "发票号"),
    "invoice_date": ("开票日期",),
    "taxpayer_no": ("纳税人识别号",),
}


@dataclass
class SourceIssue:
    source_file: str
    sheet: str
    row: int | None
    level: str
    message: str


@dataclass
class ParsedSource:
    contracts: dict[str, dict[str, Any]] = field(default_factory=dict)
    packing_items: list[dict[str, Any]] = field(default_factory=list)
    sales_items: list[dict[str, Any]] = field(default_factory=list)
    invoice_summary_items: list[dict[str, Any]] = field(default_factory=list)
    skipped_rows: list[dict[str, Any]] = field(default_factory=list)
    workbook_mismatches: list[dict[str, Any]] = field(default_factory=list)
    skipped_workbooks: list[dict[str, Any]] = field(default_factory=list)
    issues: list[SourceIssue] = field(default_factory=list)
    workbook_count: int = 0
    exp_workbook_count: int = 0


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).replace("\u3000", " ").strip()
    return re.sub(r"\s+", " ", text)


def normalize_contract_no(value: Any) -> str:
    text = normalize_text(value).upper()
    match = EXP_RE.search(text)
    if match:
        return f"EXP{match.group(1)}"
    compact = text.replace(" ", "")
    match = EXP_RE.search(compact)
    return f"EXP{match.group(1)}" if match else ""


def normalize_purchase_contract_no(value: Any) -> str:
    text = normalize_text(value).upper()
    match = CG_RE.search(text)
    if match:
        return f"CG{match.group(1)}"
    compact = text.replace(" ", "")
    match = CG_RE.search(compact)
    return f"CG{match.group(1)}" if match else ""


def to_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value).replace(",", "").replace("$", "").replace("￥", "").replace("¥", "")
    if not text or text.lower() == "nan":
        return None
    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        return None


def to_int(value: Any) -> int | None:
    number = to_number(value)
    if number is None:
        return None
    return int(round(number))


def to_iso_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = normalize_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%Y%m%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    month_match = re.search(r"([A-Za-z]{3,9})\.?\s*(\d{1,2}),?\s*(\d{4})", text)
    if month_match:
        month_name, day, year = month_match.groups()
        for pattern in ("%b", "%B"):
            try:
                month = datetime.strptime(month_name[:3], "%b").month if pattern == "%b" else datetime.strptime(month_name, "%B").month
                return date(int(year), month, int(day)).isoformat()
            except ValueError:
                continue
    return text


def excel_bool(value: Any) -> bool | None:
    text = normalize_text(value).lower()
    if not text:
        return None
    if text in ("是", "yes", "y", "true", "1"):
        return True
    if text in ("否", "no", "n", "false", "0"):
        return False
    return None


def row_values(row: Iterable[Any]) -> list[Any]:
    return [cell.value for cell in row]


def find_exp_in_values(values: Iterable[Any]) -> str:
    for value in values:
        contract_no = normalize_contract_no(value)
        if contract_no:
            return contract_no
    return ""


def is_invalid_source(path: Path) -> bool:
    name = path.name.lower()
    return any(marker.lower() in name for marker in INVALID_FILENAME_MARKERS)


def is_relevant_source(path: Path) -> bool:
    name = path.name
    if normalize_contract_no(name):
        return True
    return any(marker in name for marker in ("出货汇总", "出货清单", "装货", "清单", "报关记录"))


def header_score(values: list[Any], aliases: dict[str, tuple[str, ...]]) -> tuple[int, dict[str, int]]:
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(values):
        header = normalize_text(raw).lower().replace(" ", "")
        if not header:
            continue
        for key, names in aliases.items():
            if key in mapping:
                continue
            for name in names:
                needle = name.lower().replace(" ", "")
                if needle and needle in header:
                    mapping[key] = idx
                    break
    return len(mapping), mapping


def find_header(ws: Any, aliases: dict[str, tuple[str, ...]], required: tuple[str, ...]) -> tuple[int | None, dict[str, int]]:
    best_row = None
    best_mapping: dict[str, int] = {}
    best_score = 0
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 35)), start=1):
        values = row_values(row)
        score, mapping = header_score(values, aliases)
        if all(key in mapping for key in required) and score > best_score:
            best_row = row_index
            best_mapping = mapping
            best_score = score
    return best_row, best_mapping


def get_cell(values: list[Any], mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def disambiguate_sales_mapping(ws: Any, header_row: int, mapping: dict[str, int]) -> dict[str, int]:
    next_mapping = dict(mapping)
    product_idx = next_mapping.get("product_name")
    spec_idx = next_mapping.get("specification")
    if product_idx is None or spec_idx != product_idx:
        return next_mapping
    headers = row_values(next(ws.iter_rows(min_row=header_row, max_row=header_row)))
    for idx, raw in enumerate(headers):
        header = normalize_text(raw).lower().replace(" ", "")
        if idx == product_idx:
            continue
        if "包装" in header or "package" in header:
            next_mapping["specification"] = idx
            return next_mapping
    return next_mapping


def contract_from_workbook(ws: Any) -> dict[str, Any]:
    all_values: list[Any] = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
        all_values.extend(row)
    contract_no = find_exp_in_values(all_values)
    buyer = ""
    signed_at = None
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
        values = [normalize_text(v) for v in row]
        line = " ".join(v for v in values if v)
        if not buyer:
            buyer_match = re.search(r"Buyer[:：]?\s*(.+)", line, re.IGNORECASE)
            if buyer_match:
                buyer = buyer_match.group(1).strip()
        if not signed_at:
            date_match = re.search(r"Date[:：]?\s*(.+)", line, re.IGNORECASE)
            if date_match:
                signed_at = to_iso_date(date_match.group(1).strip())
    return {"contract_no": contract_no, "buyer": buyer, "signed_at": signed_at}


def upsert_contract(parsed: ParsedSource, contract_no: str, source_path: Path, patch: dict[str, Any]) -> None:
    if not contract_no:
        return
    contract = parsed.contracts.setdefault(
        contract_no,
        {
            "contract_no": contract_no,
            "buyer": "",
            "signed_at": None,
            "shipped_at": None,
            "port": "",
            "customs_broker": "",
            "container_labels": [],
            "stores": [],
            "purchase_contract_nos": [],
            "invoice_nos": [],
            "source_files": [],
            "packing_item_count": 0,
            "sales_item_count": 0,
            "invoice_summary_item_count": 0,
            "total_boxes": None,
            "gross_weight": None,
            "net_weight": None,
            "volume": None,
            "total_usd": None,
            "purchase_cost_rmb": None,
        },
    )
    rel_source = str(source_path)
    if rel_source not in contract["source_files"]:
        contract["source_files"].append(rel_source)
    for key, value in patch.items():
        if value in (None, ""):
            continue
        if key in ("container_labels", "stores", "purchase_contract_nos", "invoice_nos"):
            current = contract[key]
            values = value if isinstance(value, list) else [value]
            for item in values:
                if item and item not in current:
                    current.append(item)
        elif key in ("total_boxes", "gross_weight", "net_weight", "volume", "total_usd", "purchase_cost_rmb"):
            old = contract.get(key)
            contract[key] = float(value) + float(old or 0)
        elif not contract.get(key):
            contract[key] = value


def parse_packing_sheet(parsed: ParsedSource, ws: Any, source_path: Path, workbook_contract_no: str) -> None:
    sheet_name = ws.title.replace(" ", "")
    if sheet_name in ("合同", "合 同", "发票", "发 票", "发票汇总"):
        return
    header_row, mapping = find_header(ws, PACKING_HEADER_ALIASES, ("product_name",))
    if not header_row:
        return
    has_data_shape = any(
        key in mapping
        for key in (
            "contract_no",
            "boxes",
            "gross_weight",
            "net_weight",
            "volume",
            "container_label",
            "purchase_contract_no",
            "invoice_no",
            "hs_code",
        )
    )
    if not has_data_shape:
        return

    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row), start=header_row + 1):
        values = row_values(row)
        product_name = normalize_text(get_cell(values, mapping, "product_name"))
        if not product_name or product_name.lower() == "nan":
            continue
        first_cell = normalize_text(values[0] if values else "")
        if any(marker in first_cell.upper() for marker in ("TOTAL", "TOTOL")) or "合计" in first_cell or "总计" in first_cell:
            break
        contract_no = normalize_contract_no(get_cell(values, mapping, "contract_no")) or workbook_contract_no
        if workbook_contract_no and contract_no and contract_no != workbook_contract_no and "装货" in ws.title:
            parsed.skipped_rows.append(
                {
                    "source_file": str(source_path),
                    "sheet": ws.title,
                    "row": row_idx,
                    "reason": "装货 sheet 合同号与文件合同号不一致",
                    "workbook_contract_no": workbook_contract_no,
                    "row_contract_no": contract_no,
                    "product_name": product_name,
                }
            )
            continue
        if not contract_no:
            parsed.issues.append(SourceIssue(str(source_path), ws.title, row_idx, "warn", "装箱行缺少合同号，已跳过"))
            continue

        purchase_contract_no = normalize_purchase_contract_no(get_cell(values, mapping, "purchase_contract_no"))
        item = {
            "contract_no": contract_no,
            "source_file": str(source_path),
            "sheet": ws.title,
            "row": row_idx,
            "product_name": product_name,
            "hs_code": normalize_text(get_cell(values, mapping, "hs_code")),
            "supplement": normalize_text(get_cell(values, mapping, "supplement")),
            "store": normalize_text(get_cell(values, mapping, "store")),
            "port": normalize_text(get_cell(values, mapping, "port")),
            "quantity": to_number(get_cell(values, mapping, "quantity")),
            "unit": normalize_text(get_cell(values, mapping, "unit")),
            "manufacturer": normalize_text(get_cell(values, mapping, "manufacturer")),
            "specification": normalize_text(get_cell(values, mapping, "specification")),
            "boxes": to_int(get_cell(values, mapping, "boxes")),
            "gross_weight": to_number(get_cell(values, mapping, "gross_weight")),
            "net_weight": to_number(get_cell(values, mapping, "net_weight")),
            "volume": to_number(get_cell(values, mapping, "volume")),
            "container_label": normalize_text(get_cell(values, mapping, "container_label")),
            "shipped_at": to_iso_date(get_cell(values, mapping, "shipped_at")),
            "customs_broker": normalize_text(get_cell(values, mapping, "customs_broker")),
            "is_fumigated": excel_bool(get_cell(values, mapping, "is_fumigated")),
            "purchase_contract_no": purchase_contract_no,
            "purchase_cost": to_number(get_cell(values, mapping, "purchase_cost")),
            "invoice_no": normalize_text(get_cell(values, mapping, "invoice_no")),
            "unit_price": to_number(get_cell(values, mapping, "unit_price")),
            "total_price": to_number(get_cell(values, mapping, "total_price")),
            "note": normalize_text(get_cell(values, mapping, "note")),
        }
        parsed.packing_items.append(item)
        upsert_contract(
            parsed,
            contract_no,
            source_path,
            {
                "shipped_at": item["shipped_at"],
                "port": item["port"],
                "customs_broker": item["customs_broker"],
                "container_labels": item["container_label"],
                "stores": item["store"],
                "purchase_contract_nos": purchase_contract_no,
                "invoice_nos": item["invoice_no"],
                "total_boxes": item["boxes"],
                "gross_weight": item["gross_weight"],
                "net_weight": item["net_weight"],
                "volume": item["volume"],
                "purchase_cost_rmb": item["purchase_cost"],
            },
        )


def parse_sales_sheet(parsed: ParsedSource, ws: Any, source_path: Path, workbook_contract_no: str) -> None:
    if not workbook_contract_no:
        return
    if ws.title.replace(" ", "") not in ("合同", "合 同", "发票", "发 票"):
        return
    header_row, mapping = find_header(ws, SALES_HEADER_ALIASES, ("product_name", "quantity"))
    if not header_row:
        return
    mapping = disambiguate_sales_mapping(ws, header_row, mapping)
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row), start=header_row + 1):
        values = row_values(row)
        product_name = normalize_text(get_cell(values, mapping, "product_name"))
        if not product_name or product_name.lower() == "nan":
            continue
        first_cell = normalize_text(values[0] if values else "")
        if any(marker in first_cell.upper() for marker in ("TOTAL", "TOTOL")) or "合计" in first_cell or "总计" in first_cell:
            break
        item = {
            "contract_no": workbook_contract_no,
            "source_file": str(source_path),
            "sheet": ws.title,
            "row": row_idx,
            "product_name": product_name,
            "specification": normalize_text(get_cell(values, mapping, "specification")),
            "quantity": to_number(get_cell(values, mapping, "quantity")),
            "unit": normalize_text(get_cell(values, mapping, "unit")),
            "unit_price": to_number(get_cell(values, mapping, "unit_price")),
            "total_price": to_number(get_cell(values, mapping, "total_price")),
        }
        parsed.sales_items.append(item)
        upsert_contract(
            parsed,
            workbook_contract_no,
            source_path,
            {"total_usd": item["total_price"]},
        )


def parse_invoice_summary_sheet(parsed: ParsedSource, ws: Any, source_path: Path, workbook_contract_no: str) -> None:
    if "发票汇总" not in ws.title and "开票" not in ws.title:
        return
    header_row, mapping = find_header(ws, INVOICE_HEADER_ALIASES, ("product_name",))
    if not header_row:
        return
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row), start=header_row + 1):
        values = row_values(row)
        product_name = normalize_text(get_cell(values, mapping, "product_name"))
        if not product_name or product_name.lower() == "nan":
            continue
        contract_no = workbook_contract_no or find_exp_in_values(values)
        if not contract_no:
            parsed.issues.append(SourceIssue(str(source_path), ws.title, row_idx, "warn", "发票汇总行缺少合同号，已跳过"))
            continue
        item = {
            "contract_no": contract_no,
            "source_file": str(source_path),
            "sheet": ws.title,
            "row": row_idx,
            "product_name": product_name,
            "unit": normalize_text(get_cell(values, mapping, "unit")),
            "quantity": to_number(get_cell(values, mapping, "quantity")),
            "unit_price_without_tax": to_number(get_cell(values, mapping, "unit_price_without_tax")),
            "amount_without_tax": to_number(get_cell(values, mapping, "amount_without_tax")),
            "tax_rate": normalize_text(get_cell(values, mapping, "tax_rate")),
            "tax_amount": to_number(get_cell(values, mapping, "tax_amount")),
            "amount_with_tax": to_number(get_cell(values, mapping, "amount_with_tax")),
            "declaration_no": normalize_text(get_cell(values, mapping, "declaration_no")),
            "invoice_no": normalize_text(get_cell(values, mapping, "invoice_no")),
            "invoice_date": to_iso_date(get_cell(values, mapping, "invoice_date")),
            "taxpayer_no": normalize_text(get_cell(values, mapping, "taxpayer_no")),
        }
        parsed.invoice_summary_items.append(item)
        upsert_contract(
            parsed,
            contract_no,
            source_path,
            {"invoice_nos": item["invoice_no"]},
        )


def parse_workbook(path: Path, parsed: ParsedSource, root: Path) -> None:
    if is_invalid_source(path):
        return
    rel_path = path.relative_to(root.parent) if path.is_relative_to(root.parent) else path
    if not is_relevant_source(path):
        parsed.skipped_workbooks.append(
            {
                "source_file": str(rel_path),
                "reason": "文件名不包含 EXP，且不是出货汇总/出货清单/装货清单类工作簿",
            }
        )
        return
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        parsed.issues.append(SourceIssue(str(rel_path), "", None, "error", f"Excel 打开失败: {exc}"))
        return

    parsed.workbook_count += 1
    filename_contract_no = normalize_contract_no(path.name)
    workbook_contract_no = filename_contract_no
    contract_patch: dict[str, Any] = {}
    for ws in workbook.worksheets:
        if ws.title.replace(" ", "") in ("合同", "合 同"):
            contract_patch = contract_from_workbook(ws)
            workbook_contract_no = contract_patch.get("contract_no") or workbook_contract_no
            break
    if filename_contract_no:
        parsed.exp_workbook_count += 1
    if filename_contract_no and workbook_contract_no and filename_contract_no != workbook_contract_no:
        parsed.workbook_mismatches.append(
            {
                "source_file": str(rel_path),
                "filename_contract_no": filename_contract_no,
                "sheet_contract_no": workbook_contract_no,
                "decision": "按工作簿内容中的合同号归集",
            }
        )
    if workbook_contract_no:
        upsert_contract(parsed, workbook_contract_no, rel_path, contract_patch)

    for ws in workbook.worksheets:
        parse_packing_sheet(parsed, ws, rel_path, workbook_contract_no)
        parse_sales_sheet(parsed, ws, rel_path, workbook_contract_no)
        parse_invoice_summary_sheet(parsed, ws, rel_path, workbook_contract_no)


def load_db_contracts(db_path: Path) -> set[str]:
    if not db_path.exists():
        return set()
    with sqlite3.connect(str(db_path)) as conn:
        rows = conn.execute("SELECT contractNo FROM sales_contracts").fetchall()
    return {normalize_text(row[0]).upper() for row in rows if row[0]}


def load_db_contract_snapshot(db_path: Path) -> dict[str, dict[str, Any]]:
    if not db_path.exists():
        return {}
    query = """
        SELECT
            sc.contractNo,
            sc.totalAmount,
            sc.totalBoxes,
            sc.grossWeight,
            sc.netWeight,
            sc.volume,
            COUNT(DISTINCT si.id) AS salesItemCount,
            COUNT(DISTINCT pi.id) AS packingItemCount
        FROM sales_contracts sc
        LEFT JOIN sales_items si ON si.salesContractId = sc.id
        LEFT JOIN packing_items pi ON pi.salesContractId = sc.id
        GROUP BY sc.id
    """
    with sqlite3.connect(str(db_path)) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute(query).fetchall()
    return {
        normalize_text(row["contractNo"]).upper(): {
            "contract_no": normalize_text(row["contractNo"]).upper(),
            "db_total_amount": row["totalAmount"],
            "db_total_boxes": row["totalBoxes"],
            "db_gross_weight": row["grossWeight"],
            "db_net_weight": row["netWeight"],
            "db_volume": row["volume"],
            "db_sales_item_count": row["salesItemCount"],
            "db_packing_item_count": row["packingItemCount"],
        }
        for row in rows
        if row["contractNo"]
    }


def packing_sheet_rank(sheet_name: str) -> int:
    normalized = sheet_name.replace(" ", "")
    if "装货" in normalized:
        return 0
    if "出货清单" in normalized or "装货清单" in normalized:
        return 1
    if "汇总" in normalized:
        return 2
    if "装箱清单" in normalized:
        return 3
    if "箱" in normalized:
        return 4
    return 9


def is_mixed_store_text(value: Any) -> bool:
    text = normalize_text(value)
    return any(separator in text for separator in ("、", "，", ",", ";", "；"))


def is_mixed_packing_row(row: dict[str, Any]) -> bool:
    return is_mixed_store_text(row.get("store")) or "混合" in normalize_text(row.get("port"))


def packing_split_identity(row: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        normalize_text(row.get("contract_no")),
        normalize_text(row.get("product_name")),
        normalize_text(row.get("unit")),
        normalize_text(row.get("manufacturer")),
        normalize_text(row.get("specification")),
    )


def packing_source_ref(row: dict[str, Any]) -> str:
    return f"{row.get('source_file', '')}#{row.get('sheet', '')}:{row.get('row', '')}"


def numbers_match(left: Any, right: Any, tolerance: float = 0.001) -> bool:
    if left in (None, "") and right in (None, ""):
        return True
    if left in (None, "") or right in (None, ""):
        return False
    return abs(float(left) - float(right)) <= tolerance


def rows_sum_to_target(target: dict[str, Any], rows: list[dict[str, Any]]) -> bool:
    if len(rows) < 2:
        return False
    number_fields = ("quantity", "boxes", "gross_weight", "net_weight", "volume")
    for field in number_fields:
        target_value = target.get(field)
        if target_value in (None, ""):
            continue
        row_values_for_field = [row.get(field) for row in rows if row.get(field) not in (None, "")]
        if len(row_values_for_field) != len(rows):
            return False
        tolerance = 0.01 if field == "volume" else 0.001
        if not numbers_match(target_value, sum(float(value) for value in row_values_for_field), tolerance):
            return False
    return True


def split_rows_cover_mixed_row(target: dict[str, Any], rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    target_rank = packing_sheet_rank(target["sheet"])
    candidates = [
        row
        for row in rows
        if row is not target
        and packing_sheet_rank(row["sheet"]) > target_rank
        and packing_split_identity(row) == packing_split_identity(target)
        and row.get("store")
        and row.get("port")
    ]
    return candidates if rows_sum_to_target(target, candidates) else []


def row_has_packing_evidence(row: dict[str, Any]) -> bool:
    return any(
        row.get(key) not in (None, "")
        for key in (
            "store",
            "port",
            "quantity",
            "boxes",
            "gross_weight",
            "net_weight",
            "volume",
            "container_label",
            "shipped_at",
            "customs_broker",
        )
    )


def preferred_packing_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_contract: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        by_contract[item["contract_no"]].append(item)
    selected: list[dict[str, Any]] = []
    for contract_no, rows in by_contract.items():
        candidate_rows = [row for row in rows if row_has_packing_evidence(row)] or rows
        best_rank = min(packing_sheet_rank(row["sheet"]) for row in candidate_rows)
        contract_selected: list[dict[str, Any]] = []
        for row in rows:
            if row in candidate_rows and packing_sheet_rank(row["sheet"]) == best_rank:
                contract_selected.append({**row, "selection_rule": f"rank:{best_rank}"})
        for row in list(contract_selected):
            if not is_mixed_packing_row(row):
                continue
            split_rows = split_rows_cover_mixed_row(row, rows)
            if not split_rows:
                continue
            contract_selected.remove(row)
            for split_row in split_rows:
                contract_selected.append({**split_row, "selection_rule": f"split_covers:{packing_source_ref(row)}"})
        selected.extend(contract_selected)
    return sorted(selected, key=lambda row: (row["contract_no"], row["source_file"], row["row"]))


def preferred_sales_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected = [
        {**item, "selection_rule": "合同 sheet"}
        for item in items
        if "合" in item["sheet"].replace(" ", "")
    ]
    return sorted(selected, key=lambda row: (row["contract_no"], row["source_file"], row["row"]))


def sheet_breakdown(parsed: ParsedSource) -> list[dict[str, Any]]:
    counter = Counter()
    for item in parsed.packing_items:
        counter[("packing", item["sheet"])] += 1
    for item in parsed.sales_items:
        counter[("sales", item["sheet"])] += 1
    for item in parsed.invoice_summary_items:
        counter[("invoice_summary", item["sheet"])] += 1
    return [
        {"record_type": record_type, "sheet": sheet, "rows": rows}
        for (record_type, sheet), rows in sorted(counter.items())
    ]


def build_db_comparison(
    parsed: ParsedSource,
    db_snapshot: dict[str, dict[str, Any]],
    preferred_packing: list[dict[str, Any]],
    preferred_sales: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    source_contracts = set(parsed.contracts)
    db_contracts = set(db_snapshot)
    packing_counts = Counter(item["contract_no"] for item in preferred_packing)
    sales_counts = Counter(item["contract_no"] for item in preferred_sales)
    invoice_counts = Counter(item["contract_no"] for item in parsed.invoice_summary_items)
    rows = []
    for contract_no in sorted(source_contracts | db_contracts):
        contract = parsed.contracts.get(contract_no, {})
        db = db_snapshot.get(contract_no, {})
        rows.append(
            {
                "contract_no": contract_no,
                "source_status": "present" if contract_no in source_contracts else "missing",
                "db_status": "present" if contract_no in db_contracts else "missing",
                "source_preferred_packing_count": packing_counts[contract_no],
                "db_packing_item_count": db.get("db_packing_item_count"),
                "source_preferred_sales_count": sales_counts[contract_no],
                "db_sales_item_count": db.get("db_sales_item_count"),
                "source_invoice_summary_count": invoice_counts[contract_no],
                "source_total_boxes": contract.get("total_boxes"),
                "db_total_boxes": db.get("db_total_boxes"),
                "source_gross_weight": contract.get("gross_weight"),
                "db_gross_weight": db.get("db_gross_weight"),
                "source_net_weight": contract.get("net_weight"),
                "db_net_weight": db.get("db_net_weight"),
                "source_volume": contract.get("volume"),
                "db_volume": db.get("db_volume"),
                "source_total_usd": contract.get("total_usd"),
                "db_total_amount": db.get("db_total_amount"),
                "source_files": "; ".join(contract.get("source_files", [])),
            }
        )
    return rows


def write_json(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def summarize(parsed: ParsedSource, db_contracts: set[str]) -> dict[str, Any]:
    for contract in parsed.contracts.values():
        contract_no = contract["contract_no"]
        contract["packing_item_count"] = sum(1 for item in parsed.packing_items if item["contract_no"] == contract_no)
        contract["sales_item_count"] = sum(1 for item in parsed.sales_items if item["contract_no"] == contract_no)
        contract["invoice_summary_item_count"] = sum(1 for item in parsed.invoice_summary_items if item["contract_no"] == contract_no)

    source_contracts = set(parsed.contracts)
    return {
        "workbooks_read": parsed.workbook_count,
        "exp_workbooks_read": parsed.exp_workbook_count,
        "source_contracts": len(source_contracts),
        "source_packing_items": len(parsed.packing_items),
        "source_sales_items": len(parsed.sales_items),
        "source_invoice_summary_items": len(parsed.invoice_summary_items),
        "skipped_contaminated_rows": len(parsed.skipped_rows),
        "workbook_contract_mismatches": len(parsed.workbook_mismatches),
        "skipped_workbooks": len(parsed.skipped_workbooks),
        "issues": len(parsed.issues),
        "db_sales_contracts": len(db_contracts),
        "contracts_missing_in_db": sorted(source_contracts - db_contracts),
        "db_contracts_missing_in_source": sorted(db_contracts - source_contracts),
        "contracts_in_both": sorted(source_contracts & db_contracts),
    }


def write_report(
    out_dir: Path,
    summary: dict[str, Any],
    parsed: ParsedSource,
    db_comparison: list[dict[str, Any]],
) -> None:
    needs_backfill = [
        row for row in db_comparison
        if row["source_status"] == "present"
        and row["db_status"] == "present"
        and (
            row["source_preferred_packing_count"] != (row["db_packing_item_count"] or 0)
            or row["source_preferred_sales_count"] != (row["db_sales_item_count"] or 0)
        )
    ]
    lines = [
        "# WPS 出货源文件导入前核对报告",
        "",
        "## 范围",
        "",
        "- 源目录：`tmp/wps_11_export_list_raw/11-报关记录`",
        "- 当前脚本只做解析、比对、落盘报告，不直接改数据库。",
        "- 缺失字段保持空值；不会从文件名或相邻单元格之外编造供应商、门店、金额或日期。",
        "",
        "## 摘要",
        "",
        f"- 已读取 Excel 工作簿：{summary['workbooks_read']}",
        f"- 其中 EXP 出口合同工作簿：{summary['exp_workbooks_read']}",
        f"- 解析到出口合同：{summary['source_contracts']}",
        f"- 解析到装箱/出货行：{summary['source_packing_items']}",
        f"- 解析到合同/发票商品行：{summary['source_sales_items']}",
        f"- 解析到发票汇总行：{summary['source_invoice_summary_items']}",
        f"- 因装货 sheet 合同号串表而跳过的行：{summary['skipped_contaminated_rows']}",
        f"- 文件名合同号与内容合同号不一致：{summary['workbook_contract_mismatches']}",
        f"- 跳过的非出货/非 EXP 工作簿：{summary['skipped_workbooks']}",
        f"- 解析警告/错误：{summary['issues']}",
        "",
        "## 现库差异",
        "",
        f"- 当前数据库出口合同：{summary['db_sales_contracts']}",
        f"- 源文件有、数据库没有：{len(summary['contracts_missing_in_db'])}",
        f"- 数据库有、源文件没有：{len(summary['db_contracts_missing_in_source'])}",
        f"- 两边都有：{len(summary['contracts_in_both'])}",
        f"- 两边都有但行数不一致、需要补齐/重算：{len(needs_backfill)}",
        "",
        "### 源文件有、数据库没有",
        "",
    ]
    lines.extend(f"- {item}" for item in summary["contracts_missing_in_db"][:100])
    if not summary["contracts_missing_in_db"]:
        lines.append("- 无")
    lines.extend(["", "### 两边都有但需要补齐/重算", ""])
    if needs_backfill:
        for row in needs_backfill[:100]:
            lines.append(
                "- "
                f"{row['contract_no']}：装箱 源{row['source_preferred_packing_count']}/库{row['db_packing_item_count']}，"
                f"销售 源{row['source_preferred_sales_count']}/库{row['db_sales_item_count']}"
            )
    else:
        lines.append("- 无")
    lines.extend(["", "### 文件名与内容合同号不一致", ""])
    if parsed.workbook_mismatches:
        for item in parsed.workbook_mismatches:
            lines.append(
                f"- {item['source_file']}：文件名 {item['filename_contract_no']}，内容 {item['sheet_contract_no']}；{item['decision']}"
            )
    else:
        lines.append("- 无")
    lines.extend(["", "### 需要人工复核的解析问题", ""])
    if parsed.issues:
        for issue in parsed.issues[:100]:
            row = f":{issue.row}" if issue.row is not None else ""
            lines.append(f"- [{issue.level}] {issue.source_file} / {issue.sheet}{row}：{issue.message}")
    else:
        lines.append("- 无")
    lines.extend(
        [
            "",
            "## 下一步导入原则",
            "",
            "- 优先导入 `contracts_missing_in_db` 中的合同，再对 `contracts_in_both` 做字段补齐。",
            "- `装货` sheet 中合同号与文件名合同号不一致的行已跳过，应人工确认后再导入。",
            "- 落库前先按合同号核对 `contracts.csv`、`preferred_packing_items.csv`、`preferred_sales_items.csv`、`invoice_summary_items.csv` 与 `db_comparison.csv`。",
            "- 采购合同附件、报关单 PDF、退税联等文件应作为附件/凭证保留，不应混入商品行字段。",
        ]
    )
    (out_dir / "import_gap_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="解析 WPS 11-报关记录出货源文件并生成导入前核对报告")
    parser.add_argument("--source", default="tmp/wps_11_export_list_raw/11-报关记录")
    parser.add_argument("--db", default="backend/prisma/dev.db")
    parser.add_argument("--out", default="tmp/wps_11_export_list_raw/parsed")
    args = parser.parse_args()

    source = Path(args.source)
    db_path = Path(args.db)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    parsed = ParsedSource()
    for path in sorted(source.rglob("*.xlsx")):
        parse_workbook(path, parsed, source)

    db_snapshot = load_db_contract_snapshot(db_path)
    db_contracts = set(db_snapshot)
    selected_packing = preferred_packing_items(parsed.packing_items)
    selected_sales = preferred_sales_items(parsed.sales_items)
    comparison = build_db_comparison(parsed, db_snapshot, selected_packing, selected_sales)
    summary = summarize(parsed, db_contracts)

    contracts = sorted(parsed.contracts.values(), key=lambda item: item["contract_no"])
    write_json(out_dir / "summary.json", summary)
    write_json(out_dir / "contracts.json", contracts)
    write_json(out_dir / "issues.json", [issue.__dict__ for issue in parsed.issues])
    write_json(out_dir / "skipped_rows.json", parsed.skipped_rows)
    write_json(out_dir / "workbook_mismatches.json", parsed.workbook_mismatches)
    write_json(out_dir / "skipped_workbooks.json", parsed.skipped_workbooks)
    write_csv(out_dir / "contracts.csv", contracts)
    write_csv(out_dir / "packing_items.csv", parsed.packing_items)
    write_csv(out_dir / "preferred_packing_items.csv", selected_packing)
    write_csv(out_dir / "sales_items.csv", parsed.sales_items)
    write_csv(out_dir / "preferred_sales_items.csv", selected_sales)
    write_csv(out_dir / "invoice_summary_items.csv", parsed.invoice_summary_items)
    write_csv(out_dir / "sheet_breakdown.csv", sheet_breakdown(parsed))
    write_csv(out_dir / "db_comparison.csv", comparison)
    write_report(out_dir, summary, parsed, comparison)

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
