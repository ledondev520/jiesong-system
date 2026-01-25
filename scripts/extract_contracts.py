#!/usr/bin/env python3
"""
职责：从购销合同文档中提取商品、供应商信息
思路：
  1. 遍历捷淞汇总目录下所有 .docx 文件
  2. 解析每个合同文档，提取关键信息
  3. 去重整理成 CSV 文件

输出：商品-供应商关联表 (CSV)
"""

import os
import re
import csv
from pathlib import Path
from docx import Document

# 目标目录
SOURCE_DIR = "/Users/helena/Downloads/捷淞汇总"
OUTPUT_CSV = "/Users/helena/Cursor/jiesong_system/docs/商品供应商关联表.csv"

def extract_contract_info(doc_path):
    """
    从单个合同文档中提取信息
    返回: dict 或 None
    """
    try:
        doc = Document(doc_path)
        text = "\n".join([p.text for p in doc.paragraphs])
        
        # 提取合同编号（从文件名或段落）
        filename = os.path.basename(doc_path)
        contract_no_match = re.search(r'(CG\d{7})', filename)
        if not contract_no_match:
            contract_no_match = re.search(r'合同编号[：:]\s*(CG\d+)', text)
        contract_no = contract_no_match.group(1) if contract_no_match else ""
        
        # 提取签订日期
        signed_date = ""
        date_match = re.search(r'签订日期[：:]\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日', text)
        if date_match:
            year, month, day = date_match.groups()
            signed_date = f"{year}-{month.zfill(2)}-{day.zfill(2)}"
        
        # 提取供应商名称 - 从"乙方（供方）"后面
        supplier = ""
        supplier_match = re.search(r'乙方[（\(]供方[）\)][：:]\s*([^\n]+)', text)
        if supplier_match:
            supplier = supplier_match.group(1).strip()
            # 清理供应商名称
            supplier = re.sub(r'[\s\t]+', '', supplier)  # 移除空白
        
        # 提取商品信息（从第一个表格）
        products = []
        for table in doc.tables:
            # 检查表头是否包含产品相关列
            if len(table.rows) > 0:
                header_cells = [cell.text.strip() for cell in table.rows[0].cells]
                header_text = "".join(header_cells)
                
                # 如果是产品清单表格（表头包含"产品"或"品名"或"单位"）
                if '产品' in header_text or '品名' in header_text or ('序号' in header_text and '单位' in header_text):
                    # 找到产品和单位列的索引
                    product_col = -1
                    unit_col = -1
                    
                    for i, h in enumerate(header_cells):
                        if '产品' in h or '品名' in h or '货物' in h:
                            product_col = i
                        if h == '单位':
                            unit_col = i
                    
                    # 如果没找到产品列，假设第二列是产品
                    if product_col == -1 and len(header_cells) > 1:
                        product_col = 1
                    
                    # 遍历数据行
                    for row_idx, row in enumerate(table.rows[1:], 1):
                        cells = [cell.text.strip() for cell in row.cells]
                        
                        # 跳过合计行
                        if len(cells) > 0 and ('合计' in cells[0] or '总计' in cells[0]):
                            break
                        
                        # 提取产品名称
                        product_name = ""
                        unit = ""
                        
                        if product_col >= 0 and product_col < len(cells):
                            product_name = cells[product_col]
                        
                        if unit_col >= 0 and unit_col < len(cells):
                            unit = cells[unit_col]
                        
                        # 清理产品名称
                        if product_name and product_name != supplier:
                            # 跳过序号（纯数字）
                            if not product_name.isdigit() and len(product_name) > 1:
                                products.append({
                                    "product_name": product_name,
                                    "unit": unit
                                })
                    
                    # 找到产品清单后就退出
                    if products:
                        break
        
        # 如果没有从表格中提取到，尝试从文件名提取商品名
        if not products:
            # 文件名格式：购销合同CG2500071- 大理石桌面-圣荷西 2115.docx
            name_match = re.search(r'CG\d{7}[-_\s]*([^-]+)[-]', filename)
            if name_match:
                product_name = name_match.group(1).strip()
                if product_name and len(product_name) > 1:
                    products.append({"product_name": product_name, "unit": ""})
        
        return {
            "contract_no": contract_no,
            "supplier": supplier,
            "signed_date": signed_date,
            "products": products,
            "file_path": doc_path
        }
        
    except Exception as e:
        print(f"Error processing {doc_path}: {e}")
        return None

def main():
    """主函数"""
    # 收集所有数据
    all_records = []
    seen = set()  # 用于去重: (supplier, product_name)
    
    # 遍历所有 .docx 文件
    for root, dirs, files in os.walk(SOURCE_DIR):
        for file in files:
            if file.endswith('.docx') and '购销合同' in file:
                doc_path = os.path.join(root, file)
                print(f"Processing: {file}")
                
                info = extract_contract_info(doc_path)
                if info:
                    supplier = info['supplier']
                    if not supplier:
                        print(f"  Warning: No supplier found")
                        continue
                    
                    if not info['products']:
                        print(f"  Warning: No products found")
                        continue
                    
                    for product in info['products']:
                        key = (supplier, product['product_name'])
                        if key not in seen:
                            seen.add(key)
                            all_records.append({
                                "供应商名称": supplier,
                                "商品名称": product['product_name'],
                                "单位": product['unit'],
                                "合同编号": info['contract_no'],
                                "签订日期": info.get('signed_date', ''),
                            })
                            print(f"  + {product['product_name']} ({product['unit']}) [{info.get('signed_date', '')}]")
    
    # 按供应商名称排序
    all_records.sort(key=lambda x: x['供应商名称'])
    
    # 写入 CSV
    with open(OUTPUT_CSV, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=['供应商名称', '商品名称', '单位', '合同编号', '签订日期'])
        writer.writeheader()
        writer.writerows(all_records)
    
    print(f"\n完成！共提取 {len(all_records)} 条记录（已去重）")
    print(f"输出文件：{OUTPUT_CSV}")

if __name__ == "__main__":
    main()
