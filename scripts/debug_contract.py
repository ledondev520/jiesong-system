#!/usr/bin/env python3
"""调试：查看合同文档的实际结构"""

from docx import Document
import os

# 读取一个样本文档
doc_path = "/Users/helena/Downloads/捷淞汇总/2026年1月/购销合同CG2600001-吊灯-圣荷西625.docx"

doc = Document(doc_path)

print("=" * 50)
print("段落内容:")
print("=" * 50)
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")

print("\n" + "=" * 50)
print("表格内容:")
print("=" * 50)
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- 表格 {t_idx} ---")
    for r_idx, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"  行{r_idx}: {cells}")
