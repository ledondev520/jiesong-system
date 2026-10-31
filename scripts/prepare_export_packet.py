#!/usr/bin/env python3
"""Retired export-packet pricing proposal (kept for its read-only source readers).

Pricing is implemented only by priceFor() in backend/src/services/exportPricing.js; running this
script is refused. export_refund_rate_reference.py still imports read_dbf_records from here.
"""

from __future__ import annotations

import argparse
from contextlib import closing
import json
import math
import os
import re
import sqlite3
import sys
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


sys.path.insert(0, str(Path(__file__).resolve().parent))


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = PROJECT_ROOT / "backend/prisma/dev.db"
DEFAULT_HISTORY_ROOT = PROJECT_ROOT / "data/export-history"


def configured_path(value: str) -> Path:
    return Path(value).expanduser().resolve()


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def normalize_hs(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value))


def json_safe_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def parse_iso_day(value: str) -> str:
    datetime.strptime(value, "%Y-%m-%d")
    return value.replace("-", "")


@dataclass(frozen=True)
class RefundRecord:
    code: str
    name: str
    refund_rate: float
    start_date: str
    end_date: str
    source_version: str


def read_dbf_records(path: Path) -> Iterable[dict[str, str]]:
    raw = path.read_bytes()
    record_count = int.from_bytes(raw[4:8], "little")
    header_len = int.from_bytes(raw[8:10], "little")
    record_len = int.from_bytes(raw[10:12], "little")
    fields: list[tuple[str, int]] = []
    pos = 32
    while raw[pos] != 0x0D:
        descriptor = raw[pos : pos + 32]
        name = descriptor[:11].split(b"\0")[0].decode("ascii")
        fields.append((name, descriptor[16]))
        pos += 32
    for index in range(record_count):
        record = raw[
            header_len + index * record_len : header_len + (index + 1) * record_len
        ]
        if not record or record[0:1] == b"*":
            continue
        offset = 1
        values: dict[str, str] = {}
        for name, length in fields:
            field_raw = record[offset : offset + length]
            offset += length
            values[name] = field_raw.decode("gb18030", "replace").strip()
        yield values


def lookup_refund_rate(
    dbf_path: Path, hs_code: str, export_day: str, source_version: str
) -> RefundRecord | None:
    hs = normalize_hs(hs_code)
    candidates: list[dict[str, str]] = []
    for record in read_dbf_records(dbf_path):
        code = normalize_hs(record.get("CODE"))
        if not code or not hs.startswith(code):
            continue
        if record.get("ST_DATE", "") <= export_day <= record.get("END_DATE", ""):
            candidates.append(record)
    if not candidates:
        return None
    best = max(candidates, key=lambda row: len(normalize_hs(row.get("CODE"))))
    return RefundRecord(
        code=normalize_hs(best.get("CODE")),
        name=best.get("NAME", ""),
        refund_rate=float(best.get("TSL") or 0),
        start_date=best.get("ST_DATE", ""),
        end_date=best.get("END_DATE", ""),
        source_version=source_version,
    )


@dataclass(frozen=True)
class HistoricalQuote:
    contract_no: str
    product_name: str
    specification: str
    quantity: float
    unit_price_usd: float
    total_usd: float
    workbook: str
    workbook_mtime: float


def contract_sheet_name(sheet_names: Iterable[str]) -> str | None:
    for name in sheet_names:
        if normalize_text(name).replace("同", "同").replace(" ", "") == "合同":
            return name
    return None


def extract_contract_no(workbook_name: str, values: Iterable[Any]) -> str:
    for value in values:
        match = re.search(r"EXP\d{6}", str(value or ""), flags=re.I)
        if match:
            return match.group(0).upper()
    match = re.search(r"EXP\d{6}", workbook_name, flags=re.I)
    return match.group(0).upper() if match else ""


def scan_historical_quotes(history_root: Path) -> list[HistoricalQuote]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required for read-only history scanning") from exc

    quotes: list[HistoricalQuote] = []
    for path in sorted(history_root.rglob("外销出口合同+发票+箱单+EXP*.xlsx")):
        if path.name.startswith(".~") or "模板" in path.name:
            continue
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_name = contract_sheet_name(workbook.sheetnames)
            if not sheet_name:
                continue
            sheet = workbook[sheet_name]
            header_values = [
                cell.value
                for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8))
                for cell in row
            ]
            contract_no = extract_contract_no(path.name, header_values)
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 120)):
                values = [cell.value for cell in row[:6]]
                if len(values) < 6:
                    continue
                item, name, spec, quantity, unit_price, total = values
                if not re.fullmatch(r"\d+(?:\.0+)?", str(item or "").strip()):
                    continue
                if not name or not isinstance(unit_price, (int, float)):
                    continue
                if not isinstance(quantity, (int, float)) or quantity <= 0:
                    continue
                total_value = (
                    float(total)
                    if isinstance(total, (int, float))
                    else float(quantity) * float(unit_price)
                )
                quotes.append(
                    HistoricalQuote(
                        contract_no=contract_no,
                        product_name=str(name).strip(),
                        specification=str(spec or "").strip(),
                        quantity=float(quantity),
                        unit_price_usd=float(unit_price),
                        total_usd=total_value,
                        workbook=str(path),
                        workbook_mtime=path.stat().st_mtime,
                    )
                )
        except Exception:
            continue
    return quotes


def choose_historical_quote(
    quotes: Iterable[HistoricalQuote], product_name: str, specification: str
) -> HistoricalQuote | None:
    name_key = normalize_text(product_name)
    spec_key = normalize_text(specification)
    exact = [
        quote
        for quote in quotes
        if normalize_text(quote.product_name) == name_key
        and spec_key
        and normalize_text(quote.specification) == spec_key
    ]
    if exact:
        return max(exact, key=lambda quote: quote.workbook_mtime)
    return None


def same_name_history(
    quotes: Iterable[HistoricalQuote], product_name: str
) -> list[HistoricalQuote]:
    name_key = normalize_text(product_name)
    matches = [
        quote for quote in quotes if normalize_text(quote.product_name) == name_key
    ]
    return sorted(matches, key=lambda quote: quote.workbook_mtime, reverse=True)


def exclude_current_contract(
    quotes: Iterable[HistoricalQuote], contract_no: str
) -> list[HistoricalQuote]:
    current = normalize_text(contract_no).upper()
    return [
        quote
        for quote in quotes
        if normalize_text(quote.contract_no).upper() != current
    ]


def fetch_contract_rows(database: Path, contract_no: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    with closing(sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)) as connection:
        connection.row_factory = sqlite3.Row
        header = connection.execute(
            """
            SELECT sc.contractNo, sc.exchangeRate, sc.totalAmount, sc.containerLabel,
                   sc.shippedAt, sc.customsBroker, po.name AS port
              FROM sales_contracts sc
              LEFT JOIN ports po ON po.id = sc.portId
             WHERE sc.contractNo = ?
            """,
            (contract_no,),
        ).fetchone()
        if not header:
            raise ValueError(f"Contract not found: {contract_no}")
        rows = connection.execute(
            """
            SELECT pi.id, p.customsName, p.hsCode, p.declaration,
                   pi.quantity, pi.unit, pi.boxes, pi.grossWeight, pi.netWeight,
                   pi.volume, pi.purchaseCost, pi.specification, pi.supplement,
                   pi.manufacturer, pi.purchaseContractNo, st.name AS store
              FROM packing_items pi
              JOIN sales_contracts sc ON sc.id = pi.salesContractId
              JOIN products p ON p.id = pi.productId
              LEFT JOIN stores st ON st.id = pi.storeId
             WHERE sc.contractNo = ?
             ORDER BY pi.createdAt, pi.id
            """,
            (contract_no,),
        ).fetchall()
        return dict(header), [dict(row) for row in rows]


def read_catalog(path: Path) -> dict[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required for catalog reading") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["货品目录"]
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_text(value) for value in next(rows)]
    index = {header: position for position, header in enumerate(headers)}
    catalog: dict[str, dict[str, Any]] = {}
    for values in rows:
        name = values[index["商品品名"]]
        if not name:
            continue
        catalog[normalize_text(name)] = {
            "hsCode": values[index["HSCode"]],
            "declaration": values[index["商品要素聚合"]],
            "origin": values[index["境内货源地"]],
        }
    return catalog


def read_confirmed_goods(path: Path) -> tuple[dict[tuple[str, str], dict[str, Any]], dict[str, dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required for confirmed-workbook reading") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = next((name for name in workbook.sheetnames if normalize_text(name) == "箱单"), None)
    if not sheet_name:
        raise ValueError("Confirmed workbook missing packing-list sheet")
    sheet = workbook[sheet_name]
    exact: dict[tuple[str, str], dict[str, Any]] = {}
    by_spec_candidates: dict[str, list[dict[str, Any]]] = {}
    for row in sheet.iter_rows(min_row=9, max_row=min(sheet.max_row, 200), values_only=True):
        item, name, hs_code, specification = row[:4]
        if not isinstance(item, (int, float)) or not name:
            continue
        evidence = {
            "customsName": str(name).strip(),
            "hsCode": hs_code,
            "specification": str(specification or "").strip(),
            "declaration": row[10] if len(row) > 10 else None,
            "origin": row[11] if len(row) > 11 else None,
            "boxes": row[4] if len(row) > 4 else None,
            "grossWeight": row[5] if len(row) > 5 else None,
            "netWeight": row[6] if len(row) > 6 else None,
            "volume": row[7] if len(row) > 7 else None,
            "quantity": row[8] if len(row) > 8 else None,
            "unit": row[9] if len(row) > 9 else None,
            "source": str(path),
        }
        name_key = normalize_text(name)
        spec_key = normalize_text(specification)
        exact[(name_key, spec_key)] = evidence
        if spec_key:
            by_spec_candidates.setdefault(spec_key, []).append(evidence)
    by_spec = {
        spec: candidates[0]
        for spec, candidates in by_spec_candidates.items()
        if len(candidates) == 1
    }
    return exact, by_spec


def fetch_shipment_summary_rows(
    workbook_path: Path,
    contract_no: str,
    catalog_path: Path,
    confirmed_workbook: Path | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required for shipment-summary reading") from exc

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["出货总清单"]
    values = sheet.iter_rows(values_only=True)
    headers = [normalize_text(value) for value in next(values)]
    index = {header: position for position, header in enumerate(headers) if header}
    required = (
        "报关名", "门店", "港口", "报关数量", "单位", "厂家", "规格", "箱数",
        "毛重", "净重", "体积", "柜子编号", "出货日期", "合同号", "报关公司",
        "购销合同号", "采购金额", "商品补充信息",
    )
    missing_headers = [header for header in required if header not in index]
    if missing_headers:
        raise ValueError(f"Shipment summary missing headers: {', '.join(missing_headers)}")

    catalog = read_catalog(catalog_path)
    confirmed_exact: dict[tuple[str, str], dict[str, Any]] = {}
    confirmed_by_spec: dict[str, dict[str, Any]] = {}
    if confirmed_workbook:
        confirmed_exact, confirmed_by_spec = read_confirmed_goods(confirmed_workbook)
    matched: list[dict[str, Any]] = []
    for row_number, row in enumerate(values, start=2):
        source_contract = normalize_text(row[index["合同号"]]).upper()
        if source_contract != contract_no.upper():
            continue
        source_name = str(row[index["报关名"]] or "").strip()
        specification = str(row[index["规格"]] or "").strip()
        product = catalog.get(normalize_text(source_name), {})
        confirmed = confirmed_exact.get(
            (normalize_text(source_name), normalize_text(specification))
        ) or confirmed_by_spec.get(normalize_text(specification))
        if confirmed:
            product = confirmed
        output_name = confirmed.get("customsName") if confirmed else source_name
        evidence_conflicts: list[dict[str, Any]] = []
        if confirmed:
            source_values = {
                "quantity": row[index["报关数量"]],
                "unit": row[index["单位"]],
                "boxes": row[index["箱数"]],
                "grossWeight": row[index["毛重"]],
                "netWeight": row[index["净重"]],
                "volume": row[index["体积"]],
            }
            for field, source_value in source_values.items():
                confirmed_value = confirmed.get(field)
                if source_value in (None, "") or confirmed_value in (None, ""):
                    continue
                if isinstance(source_value, (int, float)) and isinstance(
                    confirmed_value, (int, float)
                ):
                    matches = abs(float(source_value) - float(confirmed_value)) <= 1e-6
                else:
                    matches = normalize_text(source_value) == normalize_text(confirmed_value)
                if not matches:
                    evidence_conflicts.append(
                        {
                            "field": field,
                            "shipment_summary": source_value,
                            "confirmed_workbook": confirmed_value,
                        }
                    )
        matched.append(
            {
                "id": f"shipment-summary:{row_number}",
                "sourceRow": row_number,
                "customsName": output_name,
                "hsCode": product.get("hsCode"),
                "declaration": product.get("declaration"),
                "origin": product.get("origin"),
                "quantity": row[index["报关数量"]],
                "unit": row[index["单位"]],
                "boxes": row[index["箱数"]],
                "grossWeight": row[index["毛重"]],
                "netWeight": row[index["净重"]],
                "volume": row[index["体积"]],
                "purchaseCost": row[index["采购金额"]],
                "specification": specification,
                "supplement": row[index["商品补充信息"]],
                "manufacturer": row[index["厂家"]],
                "purchaseContractNo": row[index["购销合同号"]],
                "store": row[index["门店"]],
                "port": row[index["港口"]],
                "containerLabel": row[index["柜子编号"]],
                "shippedAt": json_safe_value(row[index["出货日期"]]),
                "customsBroker": row[index["报关公司"]],
                "productEvidenceSource": (
                    "confirmed_workbook" if confirmed else "declaration_catalog"
                ),
                "evidenceConflicts": evidence_conflicts,
            }
        )
    if not matched:
        raise ValueError(f"Contract not found in shipment summary: {contract_no}")
    header = {
        "contractNo": contract_no,
        "exchangeRate": None,
        "totalAmount": None,
        "containerLabel": next((row["containerLabel"] for row in matched if row["containerLabel"]), None),
        "shippedAt": next((row["shippedAt"] for row in matched if row["shippedAt"]), None),
        "customsBroker": next((row["customsBroker"] for row in matched if row["customsBroker"]), None),
        "port": next((row["port"] for row in matched if row["port"]), None),
        "sourceWorkbook": str(workbook_path),
    }
    return header, matched


def build_proposal(args: argparse.Namespace) -> dict[str, Any]:
    """Retired: pricing lives only in backend/src/services/exportPricing.js (priceFor).

    The Python copy of the pricing rule (and the proposal built on it) was removed so that no second
    rate/markup source remains. The readers above (DBF refund lookup, history scan, catalog) stay for
    export_refund_rate_reference.py.
    """
    raise RuntimeError("报价规则只在系统服务 priceFor 中实现；本脚本不再生成报价建议")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--contract", required=True, help="EXP contract number")
    parser.add_argument("--spot-rate", required=True, type=float)
    parser.add_argument("--export-date", default=date.today().isoformat())
    parser.add_argument("--fx-buffer", type=float, help="Deprecated: must match the shared pricing file")
    parser.add_argument("--refundable-markup", type=float, help="Deprecated: must match the shared pricing file")
    parser.add_argument("--no-refund-markup", type=float, help="Deprecated: must match the shared pricing file")
    parser.add_argument(
        "--database", type=configured_path,
        default=os.environ.get("JIESONG_EXPORT_DATABASE") or str(DEFAULT_DB),
        help="SQLite path (JIESONG_EXPORT_DATABASE; default: backend/prisma/dev.db)",
    )
    parser.add_argument("--shipment-summary", type=Path)
    parser.add_argument("--declaration-catalog", type=Path)
    parser.add_argument("--confirmed-workbook", type=Path)
    parser.add_argument("--history-markup-tolerance", type=float, default=0.05)
    parser.add_argument(
        "--history-root", type=configured_path,
        default=os.environ.get("JIESONG_HISTORY_ROOT") or str(DEFAULT_HISTORY_ROOT),
        help="Local history directory (JIESONG_HISTORY_ROOT; default: data/export-history)",
    )
    parser.add_argument("--refund-dbf", type=Path, required=True)
    parser.add_argument("--refund-version", default="2026B")
    args = parser.parse_args(argv)
    # 报价规则已移到 Node 服务 backend/src/services/exportPricing.js 的 priceFor()（config/pricing.json 新结构）。
    # 本脚本的 Python 公式副本不再维护：待负责人确认生产不再使用后删除；在此之前拒绝运行，避免产出与系统不一致的价格。
    parser.error("报价规则已迁移到系统服务（priceFor），本脚本不再计算报价；请使用出口备单页面")
    if not math.isfinite(args.spot_rate) or args.spot_rate <= 0:
        parser.error("现汇必须是大于零的有限数字")
    if not math.isfinite(args.history_markup_tolerance) or not 0 <= args.history_markup_tolerance <= 1:
        parser.error("历史报价容差必须是 0 至 1 的有限数字")
    return args


if __name__ == "__main__":
    print(json.dumps(build_proposal(parse_args()), ensure_ascii=False, indent=2))
